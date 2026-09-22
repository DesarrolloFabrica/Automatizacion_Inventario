"""
flujo_lib/excel.py — lectura del Excel de rutas (RUTAS.xlsx) del run único.
---------------------------------------------------------------------------
Reúne lo que hoy hacen CLONACION_CARPETA/clone_carpeta_drive.cargar_rutas_desde_excel
y LMS_Fabrica/generar_base_rutas (leer_rutas_excel, normalizar_clasificacion):
una sola lectura con openpyxl (sin pandas) que devuelve la lista de lotes
(cliente | etiqueta | origen | destino) ya validada, o UNA ExcelInvalido que
enumera todos los problemas encontrados para corregirlos de una vez.

La columna `cliente` es opcional: cuando está vacía, la prevalidación deduce el
cliente de dónde cuelga el origen en Drive (ver flujo_lib/clasificacion.py).

Recordatorio: `destino` es la carpeta RAÍZ de destino; el flujo crea dentro
una carpeta con el nombre del origen (ver flujo_lib/destino.py).
"""

from __future__ import annotations

import re, sys, unicodedata
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from . import ROOT
from .drive import extraer_id_carpeta
from .mensajes import ErrorFlujo

try:
    from rutas_excel import resolver_rutas_excel
except ImportError:  # la raíz del repo no está en sys.path (p. ej. importado desde otra carpeta)
    sys.path.insert(0, str(ROOT))
    from rutas_excel import resolver_rutas_excel

# En el Excel hay 3 clasificaciones. En GCP `cliente` solo admite PRODUCTO/TANIA y la
# raíz real siempre es LMS_Carga: LMS_correcciones NO es carpeta raíz, se guarda como
# cliente PRODUCTO + raíz LMS_Carga (misma regla que LMS_Fabrica/generar_base_rutas.py).
CLASIFICACIONES = {
    "PRODUCTO": ("PRODUCTO", "LMS_Carga"),
    "TANIA": ("TANIA", "LMS_Carga"),
    "LMS_CORRECCIONES": ("PRODUCTO", "LMS_Carga"),
}
_ALIASES_CORRECCIONES = {
    "LMS_CORRECCIONES",
    "LMS_CORRECCION",
    "LMSCORRECIONES",
    "LMSCORRECION",
    "LMSCORRECIO",
    "LMS_CORRECC",
    "CORRECCIONES",
    "CORRECCION",
}
# Nombre canónico de columna -> nombres aceptados en la cabecera (minúsculas, sin espacios).
_ALIAS_COLUMNAS = {
    "cliente": ("cliente", "tipo", "clasificacion", "clasificación"),
    "etiqueta": ("etiqueta", "programa"),
    "origen": ("origen",),
    "destino": ("destino",),
}
_OBLIGATORIAS = ("origen", "destino")  # `cliente` es opcional: se detecta desde Drive
_FILAS_CABECERA = 10  # la cabecera debe estar en las primeras N filas
_ACCION_CORREGIR = "Corrige el Excel y vuelve a ejecutar."
_CLIENTES_VALIDOS = "PRODUCTO, TANIA o LMS_CORRECCIONES"


# ---------------------------------------------------------------------------
# Clasificación (columna cliente)
# ---------------------------------------------------------------------------
def norm_text(value: object) -> str:
    """Normaliza texto: quita acentos, mayúsculas y deja solo A-Z0-9 unidos por _."""
    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"[^A-Z0-9]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_")


def normalizar_clasificacion(valor: object) -> str | None:
    """
    Traduce la columna cliente del Excel a PRODUCTO / TANIA / LMS_CORRECCIONES.
    Acepta alias de correcciones (LMS_correccion, Correcciones, LMSCORRECIONES...).
    Devuelve None si el valor está vacío o no es válido.
    """
    texto = norm_text(valor)
    if not texto:
        return None
    if (
        texto in _ALIASES_CORRECCIONES
        or texto.startswith("LMS_CORRECC")
        or texto.startswith("LMSCORREC")
    ):
        return "LMS_CORRECCIONES"
    if texto in ("PRODUCTO", "TANIA"):
        return texto
    return None


# ---------------------------------------------------------------------------
# Lote y errores
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Lote:
    """Una fila válida del Excel: qué se clona (origen) y bajo qué raíz (destino)."""

    fila: int
    etiqueta: str
    cliente_excel: str
    clasificacion: str
    cliente_gcp: str
    raiz_gcp: str
    origen_id: str
    destino_raiz_id: str
    origen_raw: str
    destino_raw: str
    escuela_gcp: str = ""  # se completa al detectar la clasificación en Drive

    @property
    def sin_clasificar(self) -> bool:
        """True si el Excel no trajo cliente y hay que deducirlo de Drive."""
        return not self.clasificacion

    @property
    def clave(self) -> str:
        """Identidad del lote para el estado de la corrida (origen + raíz de destino)."""
        return f"{self.origen_id}|{self.destino_raiz_id}"


class ExcelInvalido(ErrorFlujo):
    """El Excel de rutas no se pudo leer o tiene errores; paso="excel" por defecto."""

    def __init__(
        self,
        motivo: str,
        accion: str = "",
        *,
        paso: str = "excel",
        detalle: str = "",
        contexto: str = "",
    ):
        super().__init__(motivo, accion, paso=paso or "excel", detalle=detalle, contexto=contexto)


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------
def _texto(celda: object) -> str:
    """Valor de celda como texto sin espacios; None -> ""."""
    return "" if celda is None else str(celda).strip()


def _celda(fila: tuple, indice: int | None) -> object:
    """Celda por índice tolerando filas cortas o columnas ausentes."""
    if indice is None or indice >= len(fila):
        return None
    return fila[indice]


def _abrir_libro(ruta: Path):
    """load_workbook en modo solo lectura con errores traducidos a ExcelInvalido."""
    nombre = ruta.name
    try:
        return openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except PermissionError as e:
        raise ExcelInvalido(
            f"El Excel {nombre} está abierto o bloqueado.",
            "Ciérralo y vuelve a ejecutar.",
            detalle=repr(e),
            contexto=nombre,
        ) from e
    except FileNotFoundError as e:
        raise ExcelInvalido(
            f"No se encontró el Excel {ruta}.",
            "Revisa la ruta indicada.",
            detalle=repr(e),
            contexto=nombre,
        ) from e
    except Exception as e:  # .xls, .csv, archivo dañado, etc.
        raise ExcelInvalido(
            f"No se pudo leer el Excel {nombre}.",
            "Revisa que sea un archivo .xlsx válido y que no esté dañado.",
            detalle=repr(e),
            contexto=nombre,
        ) from e


def _mapear_columnas(fila: tuple) -> dict[str, int | None]:
    """Índice de cada columna canónica en una fila de cabecera (None si no está)."""
    posiciones: dict[str, int] = {}
    for i, celda in enumerate(fila):
        nombre = _texto(celda).lower()
        if nombre and nombre not in posiciones:
            posiciones[nombre] = i
    columnas: dict[str, int | None] = {}
    for canonica, alias in _ALIAS_COLUMNAS.items():
        columnas[canonica] = next((posiciones[a] for a in alias if a in posiciones), None)
    return columnas


def _es_cabecera(fila: tuple) -> bool:
    valores = {_texto(c).lower() for c in fila}
    return "destino" in valores or "cliente" in valores


def _localizar_cabecera(libro, nombre: str):
    """
    Busca la fila de cabecera en las primeras filas de cada hoja (la activa primero).
    Devuelve (hoja, número de fila de la cabecera, columnas) o lanza ExcelInvalido.
    """
    hojas = list(libro.worksheets)
    activa = libro.active
    if activa in hojas:
        hojas.remove(activa)
        hojas.insert(0, activa)
    for hoja in hojas:
        # Materializar estas pocas filas evita dejar vivo el generador de una hoja
        # read-only si la validación lanza una excepción dentro del bucle (en
        # Windows ese generador puede retener el .xlsx aun después de libro.close()).
        filas = list(hoja.iter_rows(min_row=1, max_row=_FILAS_CABECERA, values_only=True))
        for numero, fila in enumerate(filas, start=1):
            if fila and _es_cabecera(fila):
                columnas = _mapear_columnas(fila)
                faltan = [c for c in _OBLIGATORIAS if columnas[c] is None]
                if faltan:
                    raise ExcelInvalido(
                        f"Al Excel {nombre} le falta la columna {', '.join(faltan)} "
                        f"en la cabecera (fila {numero} de la hoja «{hoja.title}»).",
                        "Agrega esa columna en la fila de cabecera y vuelve a ejecutar.",
                        contexto=nombre,
                        detalle=f"cabecera={fila!r}",
                    )
                return hoja, numero, columnas
    raise ExcelInvalido(
        f"No se encontró la fila de cabecera (etiqueta | origen | destino, y cliente si se indica) "
        f"en las primeras {_FILAS_CABECERA} filas del Excel {nombre}.",
        "Revisa que la hoja tenga esa cabecera y vuelve a ejecutar.",
        contexto=nombre,
    )


def _leer_fila(numero: int, fila: tuple, columnas: dict[str, int | None]) -> tuple[Lote | None, list[str]]:
    """Convierte una fila en Lote; devuelve (lote o None, lista de fallas de la fila)."""
    fallas: list[str] = []
    cliente_raw = _texto(_celda(fila, columnas["cliente"]))
    etiqueta = _texto(_celda(fila, columnas["etiqueta"])) or f"Fila {numero}"
    origen_raw = _texto(_celda(fila, columnas["origen"]))
    destino_raw = _texto(_celda(fila, columnas["destino"]))

    # `cliente` vacío no es un error: se deduce de dónde cuelga el origen en Drive
    # (ver flujo_lib/clasificacion.py). Solo se rechaza un valor escrito que no existe.
    clasificacion = normalizar_clasificacion(cliente_raw)
    if cliente_raw and clasificacion is None:
        fallas.append(f"cliente «{cliente_raw}» no reconocido (usa {_CLIENTES_VALIDOS})")

    ids: dict[str, str | None] = {}
    for columna, crudo in (("origen", origen_raw), ("destino", destino_raw)):
        if not crudo:
            fallas.append(f"enlace de {columna} vacío")
            ids[columna] = None
            continue
        ids[columna] = extraer_id_carpeta(crudo)
        if not ids[columna]:
            fallas.append(f"enlace de {columna} inválido («{crudo[:80]}»)")

    if ids["origen"] and ids["destino"] and ids["origen"] == ids["destino"]:
        fallas.append("origen y destino son la misma carpeta")

    if fallas or not ids["origen"] or not ids["destino"]:
        return None, fallas
    cliente_gcp, raiz_gcp = CLASIFICACIONES[clasificacion] if clasificacion else ("", "")
    lote = Lote(
        fila=numero,
        etiqueta=etiqueta,
        cliente_excel=cliente_raw,
        clasificacion=clasificacion or "",
        cliente_gcp=cliente_gcp,
        raiz_gcp=raiz_gcp,
        origen_id=ids["origen"],
        destino_raiz_id=ids["destino"],
        origen_raw=origen_raw,
        destino_raw=destino_raw,
    )
    return lote, []


def leer_lotes(ruta: Path) -> list[Lote]:
    """
    Lee RUTAS.xlsx y devuelve los lotes en orden de fila.

    - Cabecera: primera fila (en las 10 primeras) con "destino" o "cliente"; se
      busca primero en la hoja activa y luego en las demás.
    - Columnas: cliente|tipo|clasificacion|clasificación, etiqueta|programa,
      origen, destino. Solo origen y destino son obligatorias.
    - `cliente` es opcional: si viene vacío se deduce de dónde cuelga el origen
      en Drive (flujo_lib/clasificacion.py); si trae un valor que no existe, sí
      es error. Cuando falta, el lote queda con `sin_clasificar` en True.
    - Se saltan las filas sin nada en esas cuatro columnas. Cada fila debe tener
      enlaces de origen y destino usables y distintos; no puede repetirse el par
      origen+destino. Etiqueta vacía -> "Fila N".
    - Todos los problemas de fila se acumulan y se lanza UNA ExcelInvalido.
    """
    ruta = Path(ruta)
    nombre = ruta.name
    libro = _abrir_libro(ruta)
    try:
        hoja, fila_cabecera, columnas = _localizar_cabecera(libro, nombre)
        indices = [i for i in columnas.values() if i is not None]
        lotes: list[Lote] = []
        problemas: list[str] = []
        vistas: dict[str, int] = {}  # clave -> primera fila donde apareció
        filas = hoja.iter_rows(min_row=fila_cabecera + 1, values_only=True)
        for numero, fila in enumerate(filas, start=fila_cabecera + 1):
            if not fila or all(_texto(_celda(fila, i)) == "" for i in indices):
                continue
            lote, fallas = _leer_fila(numero, fila, columnas)
            if lote is not None:
                if lote.clave in vistas:
                    fallas.append(f"repite el origen y destino de la fila {vistas[lote.clave]}")
                else:
                    vistas[lote.clave] = numero
                    lotes.append(lote)
            if fallas:
                problemas.append(f"Fila {numero}: {'; '.join(fallas)}.")
    finally:
        libro.close()

    if problemas:
        raise ExcelInvalido(
            f"El Excel {nombre} tiene {len(problemas)} fila(s) con errores. " + " ".join(problemas),
            _ACCION_CORREGIR,
            contexto=nombre,
            detalle="\n".join(problemas),
        )
    if not lotes:
        raise ExcelInvalido(
            f"El Excel {nombre} no tiene lotes debajo de la cabecera (fila {fila_cabecera} "
            f"de la hoja «{hoja.title}»).",
            "Agrega las filas de los lotes (cliente, etiqueta, origen y destino) y vuelve a ejecutar.",
            contexto=nombre,
        )
    return lotes


def resolver_excel(explicit: Path | None = None) -> Path:
    """
    Ubica RUTAS.xlsx: --excel, variable RUTAS_XLSX, raíz del repo, CLONACION_CARPETA,
    LMS_Fabrica o la carpeta actual. Si no aparece lanza ExcelInvalido.
    """
    try:
        return resolver_rutas_excel(
            explicit, ROOT, ROOT / "CLONACION_CARPETA", ROOT / "LMS_Fabrica"
        )
    except FileNotFoundError as e:
        if explicit is not None:
            raise ExcelInvalido(
                f"No se encontró el Excel {Path(explicit)}.",
                "Revisa la ruta indicada en --excel.",
                detalle=repr(e),
                contexto=str(explicit),
            ) from e
        raise ExcelInvalido(
            "No se encontró el archivo RUTAS.xlsx.",
            'Indica la ruta con --excel "<RUTA>\\RUTAS.xlsx" y vuelve a ejecutar.',
            detalle=repr(e),
        ) from e
