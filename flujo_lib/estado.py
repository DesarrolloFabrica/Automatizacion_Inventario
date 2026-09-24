"""
flujo_lib/estado.py — estado por lote de la corrida (para reanudar sin duplicar).
-----------------------------------------------------------------------------
Guarda en corridas/<excel>.estado.json qué paso va en qué estado para cada lote
del Excel (destino, clonación, formato, verificación, carga) y el historial de
corridas. Así una reejecución retoma donde quedó y no repite lo ya hecho.
También exporta el estado a un Excel legible para adjuntarlo al correo final.

Nada de este módulo habla con Google ni con la base de datos.
"""

from __future__ import annotations

import logging
import io, json, os, tempfile
from datetime import datetime
from pathlib import Path

from . import ROOT
from .mensajes import ErrorFlujo, traducir_excepcion
from .progreso import Avance

PASOS = ("destino", "clonacion", "formato", "verificacion", "carga")
ESTADOS = ("pendiente", "en_curso", "ok", "con_diferencias", "fallido", "omitido")
DIR_CORRIDAS = ROOT / "corridas"
VERSION = 1

# Texto que se muestra en el Excel por cada estado y color de fondo (RGB sin alfa).
ETIQUETAS_ESTADO = {
    "pendiente": "Pendiente",
    "en_curso": "En curso",
    "ok": "OK",
    "con_diferencias": "Con diferencias",
    "fallido": "Fallido",
    "omitido": "Omitido",
}
COLORES_ESTADO = {
    "ok": "C6EFCE",  # verde
    "fallido": "FFC7CE",  # rojo
    "con_diferencias": "FFEB9C",  # amarillo
    "en_curso": "FFEB9C",  # amarillo
    "pendiente": "E7E6E6",  # gris
    "omitido": "E7E6E6",  # gris
}
CABECERA_ESTADO = (
    "Etiqueta",
    "Fila",
    "Cliente",
    "Carpeta destino",
    "Enlace destino",
    "Destino",
    "Clonación",
    "Formato",
    "Verificación",
    "Carga",
    "Qué pasó",
    "Qué hacer",
    "Actualizado",
)
CABECERA_CORRIDAS = ("Id", "Inicio", "Fin", "Resultado")
_ANCHOS_ESTADO = (28, 6, 12, 34, 44, 12, 14, 12, 14, 12, 50, 50, 20)
_ANCHOS_CORRIDAS = (18, 20, 20, 22)


# Cada cuántos segundos, como mucho, se escribe el JSON por un cambio de progreso.
# `avanzar` se llama cientos de veces por corrida; sin este freno el disco (o el
# bucket) se llevaría una escritura por archivo copiado.
SEGUNDOS_ENTRE_GUARDADOS = 2.0


def _progreso_vacio() -> dict:
    return {"hechos": 0, "total": 0, "mensaje": "", "porcentaje": 0}


def _paso_vacio() -> dict:
    return {
        "estado": "pendiente",
        "fecha": None,
        "detalle": "",
        "motivo": "",
        "accion": "",
        "progreso": _progreso_vacio(),
    }


def _como_progreso(avance) -> dict:
    """
    Normaliza lo que llega a `avanzar`: un Avance, un dict con hechos/total/mensaje
    o cualquier objeto con esos atributos. Siempre salen las cuatro claves.
    """
    if hasattr(avance, "como_dict"):
        datos = dict(avance.como_dict())
    elif isinstance(avance, dict):
        datos = dict(avance)
    else:
        datos = {
            "hechos": getattr(avance, "hechos", 0),
            "total": getattr(avance, "total", 0),
            "mensaje": getattr(avance, "mensaje", ""),
        }
    hechos = int(datos.get("hechos") or 0)
    total = int(datos.get("total") or 0)
    mensaje = str(datos.get("mensaje") or "")
    porcentaje = datos.get("porcentaje")
    if porcentaje is None:
        porcentaje = Avance(hechos, total, mensaje).porcentaje()
    return {"hechos": hechos, "total": total, "mensaje": mensaje, "porcentaje": int(porcentaje)}


def _enlace_carpeta(destino_id: str | None) -> str:
    return f"https://drive.google.com/drive/folders/{destino_id}" if destino_id else ""


class EstadoCorrida:
    """
    Estado persistente de un Excel de rutas: corridas realizadas y, por lote,
    el estado de cada paso. Cada cambio se guarda de inmediato (escritura atómica).
    La única excepción es el progreso (`avanzar`), que se guarda a lo sumo cada
    SEGUNDOS_ENTRE_GUARDADOS segundos y siempre al completarse el paso.

    Con `almacen` (ver flujo_lib/almacen.py) el JSON no se escribe en el disco
    sino a través de él (una carpeta o un bucket de Cloud Storage); sin almacén
    todo funciona exactamente como siempre, con archivos locales.
    """

    def __init__(self, ruta_json: Path, datos: dict, ahora=None, almacen=None):
        self.ruta_json = Path(ruta_json)
        nombre = self.ruta_json.name
        base = nombre[: -len(".json")] if nombre.endswith(".json") else nombre
        self.ruta_excel_estado = self.ruta_json.with_name(base + ".xlsx")
        self.datos = datos
        self._ahora = ahora or datetime.now
        self.corrida_id: str | None = None
        self.almacen = almacen
        self._ultimo_guardado: datetime | None = None
        self._pendiente = False  # hay progreso en memoria que aún no se escribió

    # ------------------------------------------------------------------
    # Nombres de los objetos (valen igual para el disco y para el almacén)
    # ------------------------------------------------------------------
    @property
    def nombre_json(self) -> str:
        return self.ruta_json.name

    @property
    def nombre_excel(self) -> str:
        return self.ruta_excel_estado.name

    @property
    def ruta_visible_json(self) -> str:
        """Dónde quedó el JSON (ruta del disco o gs://…), para el log y el correo."""
        if self.almacen is not None:
            return self.almacen.ruta_visible(self.nombre_json)
        return str(self.ruta_json)

    @property
    def ruta_visible_excel(self) -> str:
        """Dónde quedó el Excel de estado (ruta del disco o gs://…)."""
        if self.almacen is not None:
            return self.almacen.ruta_visible(self.nombre_excel)
        return str(self.ruta_excel_estado)

    # ------------------------------------------------------------------
    # Apertura y guardado
    # ------------------------------------------------------------------
    @classmethod
    def abrir(
        cls, excel: Path, dir_corridas: Path = DIR_CORRIDAS, ahora=None, almacen=None
    ) -> "EstadoCorrida":
        """
        Abre (o crea) corridas/<excel.stem>.estado.json. `ahora` es un callable que
        devuelve datetime (por defecto datetime.now) para poder fijar el reloj en pruebas.
        Con `almacen` el JSON se lee y se escribe por ahí, con el mismo nombre de
        objeto (<excel.stem>.estado.json), y no se toca el disco.
        """
        excel = Path(excel)
        dir_corridas = Path(dir_corridas)
        ruta_json = dir_corridas / f"{excel.stem}.estado.json"
        reloj = ahora or datetime.now
        crudo = None
        if almacen is not None:
            datos_bytes = almacen.leer(ruta_json.name)
            if datos_bytes is not None:
                crudo = datos_bytes.decode("utf-8")
        elif ruta_json.is_file():
            try:
                crudo = ruta_json.read_text(encoding="utf-8")
            except OSError as e:
                raise traducir_excepcion(e, paso="estado", contexto=ruta_json.name) from e
        if crudo is not None:
            try:
                datos = json.loads(crudo)
            except ValueError as e:
                raise ErrorFlujo(
                    f"El archivo de estado {ruta_json.name} está dañado.",
                    "Renómbralo o bórralo y vuelve a ejecutar; se creará uno nuevo "
                    "(se perderá el registro de lo ya clonado).",
                    paso="estado",
                    detalle=repr(e),
                    contexto=str(ruta_json),
                ) from e
            except OSError as e:
                raise traducir_excepcion(e, paso="estado", contexto=ruta_json.name) from e
            if not isinstance(datos, dict) or not isinstance(datos.get("lotes"), dict):
                raise ErrorFlujo(
                    f"El archivo de estado {ruta_json.name} no tiene el formato esperado.",
                    "Renómbralo o bórralo y vuelve a ejecutar; se creará uno nuevo.",
                    paso="estado",
                    contexto=str(ruta_json),
                )
            datos.setdefault("version", VERSION)
            datos.setdefault("excel", str(excel))
            datos.setdefault("corridas", [])
            return cls(ruta_json, datos, ahora, almacen)

        marca = reloj().isoformat(timespec="seconds")
        datos = {
            "version": VERSION,
            "excel": str(excel),
            "creado": marca,
            "actualizado": marca,
            "corridas": [],
            "lotes": {},
        }
        estado = cls(ruta_json, datos, ahora, almacen)
        estado.guardar()
        return estado

    def _marca(self) -> str:
        return self._ahora().isoformat(timespec="seconds")

    def _texto_json(self) -> str:
        return json.dumps(self.datos, indent=2, ensure_ascii=False, default=str) + "\n"

    def guardar(self) -> None:
        """
        Escribe el JSON. Con almacén, se lo entrega a él; sin almacén, escritura
        atómica en el disco (temporal en la misma carpeta + os.replace).
        """
        self.datos["actualizado"] = self._marca()
        self._ultimo_guardado = self._ahora()
        self._pendiente = False
        if self.almacen is not None:
            self.almacen.escribir(self.nombre_json, self._texto_json().encode("utf-8"))
            return
        self.ruta_json.parent.mkdir(parents=True, exist_ok=True)
        fd, temporal = tempfile.mkstemp(
            prefix=self.ruta_json.name + ".", suffix=".tmp", dir=str(self.ruta_json.parent)
        )
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as f:
                f.write(self._texto_json())
            os.replace(temporal, self.ruta_json)
        except BaseException:
            try:
                os.unlink(temporal)
            except OSError:
                pass
            raise

    # ------------------------------------------------------------------
    # Corridas
    # ------------------------------------------------------------------
    def iniciar_corrida(self, argumentos: dict) -> str:
        """Registra una corrida nueva (id = fecha y hora) y devuelve su id."""
        base = self._ahora().strftime("%Y%m%d_%H%M%S")
        existentes = {c.get("id") for c in self.datos["corridas"]}
        id_corrida, n = base, 2
        while id_corrida in existentes:  # dos corridas en el mismo segundo (raro)
            id_corrida, n = f"{base}_{n}", n + 1
        self.datos["corridas"].append(
            {
                "id": id_corrida,
                "inicio": self._marca(),
                "fin": None,
                "resultado": None,
                "argumentos": dict(argumentos or {}),
            }
        )
        self.corrida_id = id_corrida
        self.guardar()
        return id_corrida

    def _corrida_actual(self) -> dict:
        corridas = self.datos["corridas"]
        if self.corrida_id is not None:
            for c in corridas:
                if c.get("id") == self.corrida_id:
                    return c
        abiertas = [c for c in corridas if c.get("fin") is None]
        if abiertas:
            return abiertas[-1]
        raise ValueError("No hay ninguna corrida iniciada; llama antes a iniciar_corrida().")

    def cerrar_corrida(self, resultado: str) -> None:
        """Cierra la corrida en curso con un resultado libre ("ok", "fallido", ...)."""
        corrida = self._corrida_actual()
        corrida["fin"] = self._marca()
        corrida["resultado"] = resultado
        self.guardar()

    # ------------------------------------------------------------------
    # Lotes
    # ------------------------------------------------------------------
    @staticmethod
    def _clave_de(lote) -> str:
        clave = getattr(lote, "clave", None)
        if not clave:
            clave = f"{getattr(lote, 'origen_id', '')}|{getattr(lote, 'destino_raiz_id', '')}"
        return str(clave)

    def _lote(self, clave: str) -> dict:
        try:
            return self.datos["lotes"][clave]
        except KeyError:
            raise KeyError(
                f"El lote {clave!r} no está registrado; llama antes a registrar_lote()."
            ) from None

    def registrar_lote(self, lote) -> None:
        """
        Da de alta el lote (por su clave) con todos los pasos en "pendiente".
        Si ya existía, solo refresca etiqueta/fila/cliente y conserva sus pasos.
        Acepta cualquier objeto con atributos clave, etiqueta, fila, cliente_excel,
        origen_id y destino_raiz_id.
        """
        clave = self._clave_de(lote)
        etiqueta = str(getattr(lote, "etiqueta", "") or "")
        fila = getattr(lote, "fila", None)
        cliente = str(getattr(lote, "cliente_excel", "") or "")
        existente = self.datos["lotes"].get(clave)
        if existente is None:
            self.datos["lotes"][clave] = {
                "etiqueta": etiqueta,
                "fila": fila,
                "cliente": cliente,
                "origen_id": str(getattr(lote, "origen_id", "") or ""),
                "destino_raiz_id": str(getattr(lote, "destino_raiz_id", "") or ""),
                "destino_id": None,
                "destino_nombre": None,
                "pasos": {p: _paso_vacio() for p in PASOS},
                "ultimo_error": None,
            }
        else:
            existente["etiqueta"] = etiqueta
            existente["fila"] = fila
            existente["cliente"] = cliente
            pasos = existente.setdefault("pasos", {})
            for p in PASOS:  # por si el JSON viene de una versión con menos pasos
                pasos.setdefault(p, _paso_vacio())
                # …o de una anterior al progreso, que no traía esa clave.
                pasos[p].setdefault("progreso", _progreso_vacio())
            existente.setdefault("ultimo_error", None)
        self.guardar()

    def set_destino(self, clave: str, destino_id: str, destino_nombre: str) -> None:
        """Fija la carpeta destino real del lote (la creada o reutilizada dentro de la raíz)."""
        lote = self._lote(clave)
        lote["destino_id"] = destino_id
        lote["destino_nombre"] = destino_nombre
        self.guardar()

    def marcar(
        self,
        clave: str,
        paso: str,
        estado: str,
        *,
        detalle: str = "",
        motivo: str = "",
        accion: str = "",
        **extra,
    ) -> None:
        """
        Cambia el estado de un paso del lote y lo guarda. "fallido" fija ultimo_error;
        un "ok" posterior del mismo paso lo limpia. Claves extra se guardan tal cual.
        """
        if paso not in PASOS:
            raise ValueError(f"Paso desconocido {paso!r}; válidos: {', '.join(PASOS)}")
        if estado not in ESTADOS:
            raise ValueError(f"Estado desconocido {estado!r}; válidos: {', '.join(ESTADOS)}")
        lote = self._lote(clave)
        fecha = self._marca()
        registro = lote["pasos"].setdefault(paso, _paso_vacio())
        registro.update(
            {"estado": estado, "fecha": fecha, "detalle": detalle, "motivo": motivo, "accion": accion}
        )
        registro.update(extra)
        if estado == "fallido":
            lote["ultimo_error"] = {"paso": paso, "motivo": motivo, "accion": accion, "fecha": fecha}
        elif estado == "ok":
            ultimo = lote.get("ultimo_error")
            if ultimo and ultimo.get("paso") == paso:
                lote["ultimo_error"] = None
        self.guardar()

    def avanzar(self, clave: str, paso: str, avance) -> None:
        """
        Guarda el progreso del paso ({"hechos", "total", "mensaje", "porcentaje"})
        en lote["pasos"][paso]["progreso"].

        `avance` puede ser un Avance, un dict o cualquier objeto con hechos/total/
        mensaje. Esto se llama cientos de veces por corrida, así que el JSON se
        escribe como mucho una vez cada SEGUNDOS_ENTRE_GUARDADOS segundos (reloj
        `self._ahora`) y siempre en la última llamada del paso (hechos >= total,
        con total conocido). Lo que quede sin escribir se escribe con forzar_guardado().
        """
        if paso not in PASOS:
            raise ValueError(f"Paso desconocido {paso!r}; válidos: {', '.join(PASOS)}")
        lote = self._lote(clave)
        registro = lote["pasos"].setdefault(paso, _paso_vacio())
        progreso = _como_progreso(avance)
        registro["progreso"] = progreso
        self._pendiente = True
        ahora = self._ahora()
        completo = progreso["total"] > 0 and progreso["hechos"] >= progreso["total"]
        if self._ultimo_guardado is None:
            vencido = True
        else:
            vencido = (ahora - self._ultimo_guardado).total_seconds() >= SEGUNDOS_ENTRE_GUARDADOS
        if completo or vencido:
            self.guardar()

    def forzar_guardado(self) -> None:
        """Escribe ya lo que el freno de `avanzar` hubiera dejado en memoria."""
        if self._pendiente:
            self.guardar()

    def progreso(self, clave: str, paso: str) -> dict:
        """Copia del progreso del paso; los estados viejos sin esa clave dan uno vacío."""
        if paso not in PASOS:
            raise ValueError(f"Paso desconocido {paso!r}; válidos: {', '.join(PASOS)}")
        registro = self._lote(clave)["pasos"].get(paso) or {}
        return dict(registro.get("progreso") or _progreso_vacio())

    def paso(self, clave: str, paso: str) -> dict:
        """Copia del registro del paso (estado, fecha, detalle, motivo, accion, extras)."""
        if paso not in PASOS:
            raise ValueError(f"Paso desconocido {paso!r}; válidos: {', '.join(PASOS)}")
        copia = dict(self._lote(clave)["pasos"].get(paso) or _paso_vacio())
        copia["progreso"] = dict(copia.get("progreso") or _progreso_vacio())
        return copia

    def completado(self, clave: str, paso: str) -> bool:
        """True si ese paso del lote quedó en "ok" (en esta corrida o en una anterior)."""
        lote = self.datos["lotes"].get(clave)
        if lote is None:
            return False
        return (lote.get("pasos", {}).get(paso) or {}).get("estado") == "ok"

    def destino_de(self, clave: str) -> tuple[str | None, str | None]:
        """(destino_id, destino_nombre) del lote; (None, None) si aún no se resolvió."""
        lote = self.datos["lotes"].get(clave)
        if lote is None:
            return None, None
        return lote.get("destino_id"), lote.get("destino_nombre")

    # ------------------------------------------------------------------
    # Resumen y exportación
    # ------------------------------------------------------------------
    def resumen(self) -> list[dict]:
        """
        Una fila por lote, en orden de fila del Excel: clave, etiqueta, fila, cliente,
        destino_nombre, destino_id, pasos {paso: estado}, detalles {paso: detalle},
        progreso {paso: {hechos, total, mensaje, porcentaje}}, ultimo_error y
        actualizado (última fecha de cambio de algún paso).
        """
        filas = []
        for clave, lote in self.datos["lotes"].items():
            pasos = lote.get("pasos", {})
            fechas = [p.get("fecha") for p in pasos.values() if p.get("fecha")]
            filas.append(
                {
                    "clave": clave,
                    "etiqueta": lote.get("etiqueta", ""),
                    "fila": lote.get("fila"),
                    "cliente": lote.get("cliente", ""),
                    "destino_nombre": lote.get("destino_nombre"),
                    "destino_id": lote.get("destino_id"),
                    "pasos": {p: (pasos.get(p) or {}).get("estado", "pendiente") for p in PASOS},
                    "detalles": {p: (pasos.get(p) or {}).get("detalle", "") for p in PASOS},
                    # .get: los estados JSON de antes del progreso no tienen la clave.
                    "progreso": {
                        p: dict((pasos.get(p) or {}).get("progreso") or _progreso_vacio())
                        for p in PASOS
                    },
                    "ultimo_error": lote.get("ultimo_error"),
                    "actualizado": max(fechas) if fechas else None,
                }
            )
        filas.sort(key=lambda r: ((r["fila"] is None), r["fila"] or 0, r["etiqueta"]))
        return filas

    @staticmethod
    def _que_paso(fila: dict) -> tuple[str, str]:
        """(Qué pasó, Qué hacer) para el Excel: sale del último error del lote."""
        ultimo = fila.get("ultimo_error")
        if ultimo:
            return ultimo.get("motivo", ""), ultimo.get("accion", "")
        return "", ""

    def bytes_excel(self) -> bytes | None:
        """
        Contenido del Excel de estado, venga del almacén o del disco.

        Lo usa el correo para adjuntarlo. Devuelve None si aún no se ha
        exportado o no se puede leer: quedarse sin adjunto es molesto, pero
        quedarse sin correo es mucho peor.
        """
        try:
            if self.almacen is not None:
                return self.almacen.leer(self.nombre_excel)
            ruta = self.ruta_json.with_name(self.nombre_excel)
            return ruta.read_bytes() if ruta.is_file() else None
        except Exception as e:
            logging.getLogger(__name__).debug(
                "No se pudo leer el Excel de estado para adjuntarlo: %r", e
            )
            return None

    def exportar_excel(self, ruta: Path | None = None) -> Path:
        """
        Escribe el estado en un Excel con dos hojas ("Estado" y "Corridas") con colores
        por estado. Por defecto va junto al JSON (<stem>.estado.xlsx). Devuelve la ruta.
        Si el archivo está abierto en Excel, lanza ErrorFlujo pidiendo cerrarlo.

        Con almacén los bytes del xlsx se escriben por ahí (nombre <stem>.estado.xlsx,
        o el de `ruta` si se indica una) y no se toca el disco; la ruta que se
        devuelve sigue siendo la nominal (para verla tal cual está, `ruta_visible_excel`).
        """
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        ruta = Path(ruta) if ruta is not None else self.ruta_excel_estado
        rellenos = {
            estado: PatternFill("solid", fgColor=color) for estado, color in COLORES_ESTADO.items()
        }
        fuente_cabecera = Font(bold=True)
        relleno_cabecera = PatternFill("solid", fgColor="D6EAF8")
        ajuste = Alignment(wrap_text=True, vertical="top")
        centrado = Alignment(horizontal="center", vertical="top")

        def _cabecera(hoja, titulos, anchos):
            for col, (titulo, ancho) in enumerate(zip(titulos, anchos), start=1):
                c = hoja.cell(1, col, titulo)
                c.font = fuente_cabecera
                c.fill = relleno_cabecera
                c.alignment = Alignment(wrap_text=True, vertical="center")
                hoja.column_dimensions[get_column_letter(col)].width = ancho
            hoja.freeze_panes = "A2"

        wb = Workbook()
        hoja = wb.active
        hoja.title = "Estado"
        _cabecera(hoja, CABECERA_ESTADO, _ANCHOS_ESTADO)
        for i, fila in enumerate(self.resumen(), start=2):
            motivo, accion = self._que_paso(fila)
            enlace = _enlace_carpeta(fila["destino_id"])
            valores = [
                fila["etiqueta"],
                fila["fila"],
                fila["cliente"],
                fila["destino_nombre"] or "",
                enlace,
            ]
            for col, valor in enumerate(valores, start=1):
                c = hoja.cell(i, col, valor)
                c.alignment = ajuste
            if enlace:
                hoja.cell(i, 5).hyperlink = enlace
                hoja.cell(i, 5).style = "Hyperlink"
            for desplaz, p in enumerate(PASOS):
                estado = fila["pasos"][p]
                c = hoja.cell(i, 6 + desplaz, ETIQUETAS_ESTADO.get(estado, estado))
                c.fill = rellenos.get(estado, rellenos["pendiente"])
                c.alignment = centrado
                detalle = fila["detalles"].get(p, "")
                if detalle:  # el detalle técnico va como comentario, sin ensuciar la tabla
                    c.comment = Comment(str(detalle)[:2000], "flujo")
            for col, valor in ((11, motivo), (12, accion), (13, fila["actualizado"] or "")):
                c = hoja.cell(i, col, valor)
                c.alignment = ajuste

        hoja_c = wb.create_sheet("Corridas")
        _cabecera(hoja_c, CABECERA_CORRIDAS, _ANCHOS_CORRIDAS)
        for i, corrida in enumerate(self.datos.get("corridas", []), start=2):
            hoja_c.cell(i, 1, corrida.get("id"))
            hoja_c.cell(i, 2, corrida.get("inicio"))
            hoja_c.cell(i, 3, corrida.get("fin") or "")
            hoja_c.cell(i, 4, corrida.get("resultado") or "")

        if self.almacen is not None:
            buffer = io.BytesIO()
            try:
                wb.save(buffer)
            except Exception as e:
                raise traducir_excepcion(e, paso="estado", contexto=ruta.name) from e
            self.almacen.escribir(ruta.name, buffer.getvalue())
            return ruta

        ruta.parent.mkdir(parents=True, exist_ok=True)
        fd, temporal = tempfile.mkstemp(prefix=ruta.name + ".", suffix=".tmp", dir=str(ruta.parent))
        os.close(fd)
        try:
            wb.save(temporal)
            os.replace(temporal, ruta)
        except Exception as e:
            try:
                os.unlink(temporal)
            except OSError:
                pass
            raise traducir_excepcion(e, paso="estado", contexto=ruta.name) from e
        return ruta
