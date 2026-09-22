"""Pruebas de flujo_lib.estado: JSON de corrida, marcas por paso, resumen y Excel de estado."""

from __future__ import annotations

import json, tempfile, unittest
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import load_workbook

from flujo_lib.estado import (
    CABECERA_CORRIDAS,
    CABECERA_ESTADO,
    COLORES_ESTADO,
    ESTADOS,
    PASOS,
    EstadoCorrida,
)
from flujo_lib.mensajes import ErrorFlujo


class Reloj:
    """Reloj inyectable: devuelve siempre la misma hora hasta que se avanza."""

    def __init__(self, inicio: datetime):
        self.actual = inicio

    def __call__(self) -> datetime:
        return self.actual

    def avanzar(self, segundos: int) -> None:
        self.actual = self.actual + timedelta(seconds=segundos)


def _lote(fila: int = 5, etiqueta: str = "Bogotá — Diseño", origen: str = "ORIG1", raiz: str = "RAIZ1"):
    return SimpleNamespace(
        clave=f"{origen}|{raiz}",
        etiqueta=etiqueta,
        fila=fila,
        cliente_excel="PRODUCTO",
        origen_id=origen,
        destino_raiz_id=raiz,
    )


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name) / "corridas"
        self.excel = Path(self._tmp.name) / "RUTAS.xlsx"
        self.reloj = Reloj(datetime(2026, 9, 22, 8, 0, 0))
        self.estado = EstadoCorrida.abrir(self.excel, self.dir, ahora=self.reloj)

    def tearDown(self):
        self._tmp.cleanup()

    def _leer_json(self) -> dict:
        return json.loads(self.estado.ruta_json.read_text(encoding="utf-8"))


class TestAbrir(Base):
    def test_crea_json_con_estructura(self):
        self.assertTrue(self.estado.ruta_json.is_file())
        self.assertEqual(self.estado.ruta_json, self.dir / "RUTAS.estado.json")
        self.assertEqual(self.estado.ruta_excel_estado, self.dir / "RUTAS.estado.xlsx")
        datos = self._leer_json()
        self.assertEqual(
            set(datos), {"version", "excel", "creado", "actualizado", "corridas", "lotes"}
        )
        self.assertEqual(datos["version"], 1)
        self.assertEqual(datos["excel"], str(self.excel))
        self.assertEqual(datos["creado"], "2026-09-22T08:00:00")
        self.assertEqual(datos["actualizado"], "2026-09-22T08:00:00")
        self.assertEqual(datos["corridas"], [])
        self.assertEqual(datos["lotes"], {})

    def test_crea_dir_corridas_si_no_existe(self):
        anidado = Path(self._tmp.name) / "no" / "existe" / "corridas"
        self.assertFalse(anidado.exists())
        estado = EstadoCorrida.abrir(self.excel, anidado, ahora=self.reloj)
        self.assertTrue(anidado.is_dir())
        self.assertTrue(estado.ruta_json.is_file())

    def test_reabrir_recupera_todo(self):
        lote = _lote()
        self.estado.iniciar_corrida({"schema": "fabrica", "simular": False})
        self.estado.registrar_lote(lote)
        self.estado.set_destino(lote.clave, "DEST1", "Bogotá — Diseño")
        self.estado.marcar(lote.clave, "destino", "ok", detalle="creada")
        self.estado.marcar(lote.clave, "clonacion", "fallido", motivo="No se pudo.", accion="Reintenta.")
        self.estado.cerrar_corrida("fallido")

        otro = EstadoCorrida.abrir(self.excel, self.dir, ahora=self.reloj)
        self.assertEqual(otro.datos, self.estado.datos)
        self.assertTrue(otro.completado(lote.clave, "destino"))
        self.assertEqual(otro.destino_de(lote.clave), ("DEST1", "Bogotá — Diseño"))
        self.assertEqual(otro.paso(lote.clave, "clonacion")["estado"], "fallido")
        self.assertEqual(otro.datos["corridas"][0]["resultado"], "fallido")
        self.assertEqual(otro.datos["corridas"][0]["argumentos"], {"schema": "fabrica", "simular": False})

    def test_json_danado_lanza_error_flujo(self):
        self.estado.ruta_json.write_text("{esto no es json", encoding="utf-8")
        with self.assertRaises(ErrorFlujo) as ctx:
            EstadoCorrida.abrir(self.excel, self.dir, ahora=self.reloj)
        self.assertIn("está dañado", ctx.exception.motivo)

    def test_reloj_por_defecto_es_datetime_now(self):
        estado = EstadoCorrida.abrir(Path(self._tmp.name) / "Otro.xlsx", self.dir)
        # No se fija la hora: solo comprobamos que quedó una marca ISO válida.
        datetime.fromisoformat(estado.datos["creado"])


class TestCorridas(Base):
    def test_iniciar_y_cerrar(self):
        id_corrida = self.estado.iniciar_corrida({"excel": str(self.excel), "simular": True})
        self.assertEqual(id_corrida, "20260922_080000")
        corrida = self._leer_json()["corridas"][0]
        self.assertEqual(corrida["id"], id_corrida)
        self.assertEqual(corrida["inicio"], "2026-09-22T08:00:00")
        self.assertIsNone(corrida["fin"])
        self.assertIsNone(corrida["resultado"])
        self.assertEqual(corrida["argumentos"], {"excel": str(self.excel), "simular": True})

        self.reloj.avanzar(90)
        self.estado.cerrar_corrida("ok")
        corrida = self._leer_json()["corridas"][0]
        self.assertEqual(corrida["fin"], "2026-09-22T08:01:30")
        self.assertEqual(corrida["resultado"], "ok")
        self.assertEqual(self._leer_json()["actualizado"], "2026-09-22T08:01:30")

    def test_dos_corridas_en_el_mismo_segundo_tienen_ids_distintos(self):
        a = self.estado.iniciar_corrida({})
        b = self.estado.iniciar_corrida({})
        self.assertNotEqual(a, b)
        self.assertEqual(len(self.estado.datos["corridas"]), 2)

    def test_cerrar_sin_iniciar_es_error(self):
        with self.assertRaises(ValueError):
            self.estado.cerrar_corrida("ok")

    def test_argumentos_con_path_se_serializan(self):
        self.estado.iniciar_corrida({"excel": self.excel})
        self.assertEqual(self._leer_json()["corridas"][0]["argumentos"]["excel"], str(self.excel))


class TestLotes(Base):
    def test_registrar_lote_nuevo(self):
        lote = _lote()
        self.estado.registrar_lote(lote)
        registro = self._leer_json()["lotes"][lote.clave]
        self.assertEqual(registro["etiqueta"], "Bogotá — Diseño")
        self.assertEqual(registro["fila"], 5)
        self.assertEqual(registro["cliente"], "PRODUCTO")
        self.assertEqual(registro["origen_id"], "ORIG1")
        self.assertEqual(registro["destino_raiz_id"], "RAIZ1")
        self.assertIsNone(registro["destino_id"])
        self.assertIsNone(registro["destino_nombre"])
        self.assertIsNone(registro["ultimo_error"])
        self.assertEqual(tuple(registro["pasos"]), PASOS)
        for p in PASOS:
            self.assertEqual(
                registro["pasos"][p],
                {"estado": "pendiente", "fecha": None, "detalle": "", "motivo": "", "accion": ""},
            )

    def test_registrar_lote_existente_conserva_pasos(self):
        lote = _lote(fila=5, etiqueta="Vieja")
        self.estado.registrar_lote(lote)
        self.estado.set_destino(lote.clave, "DEST1", "Vieja")
        self.estado.marcar(lote.clave, "destino", "ok")
        self.estado.marcar(lote.clave, "clonacion", "ok", detalle="12 archivos")

        self.estado.registrar_lote(_lote(fila=9, etiqueta="Nueva"))
        registro = self.estado.datos["lotes"][lote.clave]
        self.assertEqual(registro["etiqueta"], "Nueva")
        self.assertEqual(registro["fila"], 9)
        self.assertEqual(registro["pasos"]["destino"]["estado"], "ok")
        self.assertEqual(registro["pasos"]["clonacion"]["detalle"], "12 archivos")
        self.assertEqual(self.estado.destino_de(lote.clave), ("DEST1", "Vieja"))
        self.assertEqual(len(self.estado.datos["lotes"]), 1)

    def test_registrar_lote_sin_clave_la_calcula(self):
        lote = SimpleNamespace(etiqueta="X", fila=2, cliente_excel="TANIA", origen_id="A", destino_raiz_id="B")
        self.estado.registrar_lote(lote)
        self.assertIn("A|B", self.estado.datos["lotes"])

    def test_set_destino_y_destino_de(self):
        lote = _lote()
        self.assertEqual(self.estado.destino_de(lote.clave), (None, None))
        self.estado.registrar_lote(lote)
        self.assertEqual(self.estado.destino_de(lote.clave), (None, None))
        self.estado.set_destino(lote.clave, "DEST1", "Bogotá — Diseño")
        self.assertEqual(self.estado.destino_de(lote.clave), ("DEST1", "Bogotá — Diseño"))
        registro = self._leer_json()["lotes"][lote.clave]
        self.assertEqual((registro["destino_id"], registro["destino_nombre"]), ("DEST1", "Bogotá — Diseño"))

    def test_set_destino_lote_desconocido(self):
        with self.assertRaises(KeyError):
            self.estado.set_destino("no|existe", "D", "N")


class TestMarcar(Base):
    def setUp(self):
        super().setUp()
        self.lote = _lote()
        self.estado.registrar_lote(self.lote)
        self.clave = self.lote.clave

    def test_cambia_estado_y_fecha(self):
        self.reloj.avanzar(10)
        self.estado.marcar(self.clave, "clonacion", "en_curso", detalle="copiando")
        paso = self.estado.paso(self.clave, "clonacion")
        self.assertEqual(paso["estado"], "en_curso")
        self.assertEqual(paso["fecha"], "2026-09-22T08:00:10")
        self.assertEqual(paso["detalle"], "copiando")
        self.assertEqual(self._leer_json()["lotes"][self.clave]["pasos"]["clonacion"]["estado"], "en_curso")
        # Los demás pasos no se tocan.
        self.assertEqual(self.estado.paso(self.clave, "destino")["estado"], "pendiente")

    def test_extra_se_guarda(self):
        self.estado.marcar(self.clave, "clonacion", "ok", archivos_copiados=7)
        self.assertEqual(self.estado.paso(self.clave, "clonacion")["archivos_copiados"], 7)

    def test_fallido_fija_ultimo_error(self):
        self.reloj.avanzar(5)
        self.estado.marcar(
            self.clave, "destino", "fallido", motivo="No se pudo crear.", accion="Pide acceso.", detalle="HttpError 403"
        )
        ultimo = self.estado.datos["lotes"][self.clave]["ultimo_error"]
        self.assertEqual(
            ultimo,
            {"paso": "destino", "motivo": "No se pudo crear.", "accion": "Pide acceso.", "fecha": "2026-09-22T08:00:05"},
        )
        self.assertNotIn("HttpError", json.dumps(ultimo))  # el detalle técnico no va al último error
        self.assertEqual(self.estado.paso(self.clave, "destino")["detalle"], "HttpError 403")

    def test_ok_posterior_limpia_ultimo_error(self):
        self.estado.marcar(self.clave, "clonacion", "fallido", motivo="m", accion="a")
        self.assertIsNotNone(self.estado.datos["lotes"][self.clave]["ultimo_error"])
        self.estado.marcar(self.clave, "clonacion", "ok", detalle="todo copiado")
        self.assertIsNone(self.estado.datos["lotes"][self.clave]["ultimo_error"])
        self.assertIsNone(self._leer_json()["lotes"][self.clave]["ultimo_error"])

    def test_ok_de_otro_paso_no_limpia_ultimo_error(self):
        self.estado.marcar(self.clave, "clonacion", "fallido", motivo="m", accion="a")
        self.estado.marcar(self.clave, "destino", "ok")
        self.assertEqual(self.estado.datos["lotes"][self.clave]["ultimo_error"]["paso"], "clonacion")

    def test_paso_invalido(self):
        with self.assertRaises(ValueError):
            self.estado.marcar(self.clave, "correo", "ok")
        with self.assertRaises(ValueError):
            self.estado.paso(self.clave, "correo")

    def test_estado_invalido(self):
        with self.assertRaises(ValueError):
            self.estado.marcar(self.clave, "destino", "listo")

    def test_lote_desconocido(self):
        with self.assertRaises(KeyError):
            self.estado.marcar("no|existe", "destino", "ok")

    def test_todos_los_estados_se_aceptan(self):
        for estado in ESTADOS:
            self.estado.marcar(self.clave, "carga", estado)
            self.assertEqual(self.estado.paso(self.clave, "carga")["estado"], estado)

    def test_completado(self):
        self.assertFalse(self.estado.completado(self.clave, "destino"))
        self.assertFalse(self.estado.completado("no|existe", "destino"))
        self.estado.marcar(self.clave, "destino", "con_diferencias")
        self.assertFalse(self.estado.completado(self.clave, "destino"))
        self.estado.marcar(self.clave, "destino", "ok")
        self.assertTrue(self.estado.completado(self.clave, "destino"))

    def test_paso_devuelve_copia(self):
        copia = self.estado.paso(self.clave, "destino")
        copia["estado"] = "ok"
        self.assertEqual(self.estado.paso(self.clave, "destino")["estado"], "pendiente")


class TestGuardar(Base):
    def test_guardar_es_atomico_y_legible(self):
        lote = _lote(etiqueta="Medellín — Ñandú «prueba»")
        self.estado.registrar_lote(lote)
        self.estado.marcar(lote.clave, "destino", "fallido", motivo="Sin permiso ñ", accion="Pide acceso")
        archivos = sorted(p.name for p in self.dir.iterdir())
        self.assertEqual(archivos, ["RUTAS.estado.json"])  # no queda ningún temporal
        texto = self.estado.ruta_json.read_text(encoding="utf-8")
        self.assertIn("Medellín — Ñandú «prueba»", texto)  # acentos sin escapar
        self.assertNotIn("\\u00", texto)
        datos = json.loads(texto)
        self.assertEqual(datos["lotes"][lote.clave]["etiqueta"], "Medellín — Ñandú «prueba»")
        self.assertTrue(texto.startswith("{\n  "))  # indent=2

    def test_guardar_actualiza_marca(self):
        self.reloj.avanzar(3600)
        self.estado.guardar()
        self.assertEqual(self._leer_json()["actualizado"], "2026-09-22T09:00:00")
        self.assertEqual(self._leer_json()["creado"], "2026-09-22T08:00:00")


class TestResumen(Base):
    def test_ordenado_por_fila(self):
        c = _lote(fila=12, etiqueta="C", origen="O3", raiz="R")
        a = _lote(fila=3, etiqueta="A", origen="O1", raiz="R")
        b = _lote(fila=7, etiqueta="B", origen="O2", raiz="R")
        for lote in (c, a, b):
            self.estado.registrar_lote(lote)
        self.estado.set_destino(a.clave, "DA", "A")
        self.estado.marcar(a.clave, "destino", "ok")
        self.estado.marcar(a.clave, "clonacion", "fallido", motivo="No copió.", accion="Reintenta.")

        resumen = self.estado.resumen()
        self.assertEqual([r["etiqueta"] for r in resumen], ["A", "B", "C"])
        self.assertEqual([r["fila"] for r in resumen], [3, 7, 12])
        fila_a = resumen[0]
        self.assertEqual(fila_a["clave"], a.clave)
        self.assertEqual(fila_a["cliente"], "PRODUCTO")
        self.assertEqual((fila_a["destino_id"], fila_a["destino_nombre"]), ("DA", "A"))
        self.assertEqual(
            fila_a["pasos"],
            {"destino": "ok", "clonacion": "fallido", "formato": "pendiente", "verificacion": "pendiente", "carga": "pendiente"},
        )
        self.assertEqual(fila_a["ultimo_error"]["motivo"], "No copió.")
        self.assertEqual(fila_a["actualizado"], "2026-09-22T08:00:00")
        self.assertIsNone(resumen[1]["ultimo_error"])
        self.assertIsNone(resumen[1]["actualizado"])

    def test_vacio(self):
        self.assertEqual(self.estado.resumen(), [])


class TestExportarExcel(Base):
    def setUp(self):
        super().setUp()
        self.estado.iniciar_corrida({"schema": "fabrica"})
        self.ok = _lote(fila=2, etiqueta="Lote OK", origen="O1", raiz="R")
        self.mal = _lote(fila=4, etiqueta="Lote con error", origen="O2", raiz="R")
        for lote in (self.ok, self.mal):
            self.estado.registrar_lote(lote)
        self.estado.set_destino(self.ok.clave, "DEST_OK", "Carpeta OK")
        self.estado.marcar(self.ok.clave, "destino", "ok", detalle="Carpeta destino creada")
        self.estado.marcar(self.ok.clave, "clonacion", "ok")
        self.estado.marcar(self.mal.clave, "destino", "fallido", motivo="No se encontró la carpeta.", accion="Revisa el enlace.")
        self.estado.marcar(self.mal.clave, "clonacion", "omitido", detalle="No se ejecutó porque falló el paso anterior")
        self.estado.marcar(self.mal.clave, "formato", "con_diferencias")
        self.reloj.avanzar(60)
        self.estado.cerrar_corrida("fallido")

    def _rgb(self, celda) -> str:
        return str(celda.fill.fgColor.rgb)[-6:]

    def test_crea_xlsx_con_hojas_y_cabeceras(self):
        ruta = self.estado.exportar_excel()
        self.assertEqual(ruta, self.estado.ruta_excel_estado)
        self.assertTrue(ruta.is_file())
        wb = load_workbook(ruta)
        self.assertEqual(wb.sheetnames, ["Estado", "Corridas"])
        hoja = wb["Estado"]
        self.assertEqual(tuple(c.value for c in hoja[1]), CABECERA_ESTADO)
        self.assertEqual(
            CABECERA_ESTADO,
            ("Etiqueta", "Fila", "Cliente", "Carpeta destino", "Enlace destino", "Destino", "Clonación",
             "Formato", "Verificación", "Carga", "Qué pasó", "Qué hacer", "Actualizado"),
        )
        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertEqual(hoja.max_row, 3)
        corridas = wb["Corridas"]
        self.assertEqual(tuple(c.value for c in corridas[1]), CABECERA_CORRIDAS)
        self.assertEqual(corridas.freeze_panes, "A2")
        self.assertEqual(
            [c.value for c in corridas[2]],
            ["20260922_080000", "2026-09-22T08:00:00", "2026-09-22T08:01:00", "fallido"],
        )

    def test_contenido_y_colores(self):
        wb = load_workbook(self.estado.exportar_excel())
        hoja = wb["Estado"]
        fila_ok, fila_mal = hoja[2], hoja[3]  # orden por fila del Excel (2 antes que 4)
        self.assertEqual(fila_ok[0].value, "Lote OK")
        self.assertEqual(fila_ok[1].value, 2)
        self.assertEqual(fila_ok[2].value, "PRODUCTO")
        self.assertEqual(fila_ok[3].value, "Carpeta OK")
        self.assertEqual(fila_ok[4].value, "https://drive.google.com/drive/folders/DEST_OK")
        self.assertEqual(fila_ok[4].hyperlink.target, "https://drive.google.com/drive/folders/DEST_OK")
        self.assertEqual(fila_ok[5].value, "OK")
        self.assertEqual(self._rgb(fila_ok[5]), COLORES_ESTADO["ok"])
        self.assertEqual(self._rgb(fila_ok[6]), "C6EFCE")
        self.assertEqual(fila_ok[7].value, "Pendiente")
        self.assertEqual(self._rgb(fila_ok[7]), COLORES_ESTADO["pendiente"])
        self.assertEqual((fila_ok[10].value, fila_ok[11].value), (None, None))
        self.assertEqual(fila_ok[12].value, "2026-09-22T08:00:00")

        self.assertEqual(fila_mal[0].value, "Lote con error")
        self.assertEqual(fila_mal[3].value, None)  # sin destino: celda vacía
        self.assertEqual(fila_mal[5].value, "Fallido")
        self.assertEqual(self._rgb(fila_mal[5]), "FFC7CE")
        self.assertEqual(fila_mal[6].value, "Omitido")
        self.assertEqual(self._rgb(fila_mal[6]), COLORES_ESTADO["omitido"])
        self.assertEqual(fila_mal[7].value, "Con diferencias")
        self.assertEqual(self._rgb(fila_mal[7]), "FFEB9C")
        self.assertEqual(fila_mal[10].value, "No se encontró la carpeta.")
        self.assertEqual(fila_mal[11].value, "Revisa el enlace.")

    def test_ruta_explicita(self):
        destino = Path(self._tmp.name) / "otra" / "estado_final.xlsx"
        ruta = self.estado.exportar_excel(destino)
        self.assertEqual(ruta, destino)
        self.assertTrue(destino.is_file())
        self.assertEqual(sorted(p.name for p in destino.parent.iterdir()), ["estado_final.xlsx"])

    def test_archivo_bloqueado_da_error_flujo(self):
        with mock.patch("openpyxl.workbook.workbook.Workbook.save", side_effect=PermissionError("bloqueado")):
            with self.assertRaises(ErrorFlujo) as ctx:
                self.estado.exportar_excel()
        self.assertIn("está abierto o protegido", ctx.exception.motivo)
        self.assertIn("RUTAS.estado.xlsx", ctx.exception.motivo)
        # No queda ningún temporal del intento fallido.
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["RUTAS.estado.json"])


if __name__ == "__main__":
    unittest.main()
