"""Pruebas de flujo_lib.excel con Excels temporales creados con openpyxl. Sin red."""

from __future__ import annotations

import tempfile, unittest
from pathlib import Path
from unittest import mock

import openpyxl

from flujo_lib import excel
from flujo_lib.excel import CLASIFICACIONES, ExcelInvalido, Lote, leer_lotes, normalizar_clasificacion, resolver_excel
from flujo_lib.mensajes import ErrorFlujo

CABECERA = ["cliente", "etiqueta", "origen", "destino"]
URL_A = "https://drive.google.com/drive/folders/1AaOrigenBogota_01?usp=sharing"
URL_B = "https://drive.google.com/drive/u/0/folders/1BbOrigenMedellin-02"
URL_OPEN = "https://drive.google.com/open?id=1CcDestinoRaiz_03"
ID_RAIZ = "1DdDestinoRaiz-04"


def _escribir_excel(carpeta: Path, filas: dict[int, list], nombre: str = "RUTAS.xlsx", hojas_previas: int = 0) -> Path:
    """
    Crea un .xlsx con las filas indicadas por número (1-based) en la hoja activa.
    hojas_previas > 0 agrega hojas vacías ANTES y deja activa la que tiene datos.
    """
    libro = openpyxl.Workbook()
    hoja = libro.active
    for _ in range(hojas_previas):
        hoja = libro.create_sheet()
    libro.active = libro.worksheets.index(hoja)
    for numero, valores in filas.items():
        for columna, valor in enumerate(valores, start=1):
            hoja.cell(row=numero, column=columna, value=valor)
    ruta = carpeta / nombre
    libro.save(ruta)
    libro.close()
    return ruta


class BaseExcel(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def excel(self, filas: dict[int, list], **kw) -> Path:
        return _escribir_excel(self.dir, filas, **kw)

    def leer_error(self, ruta: Path) -> ExcelInvalido:
        with self.assertRaises(ExcelInvalido) as cm:
            leer_lotes(ruta)
        return cm.exception


class TestNormalizarClasificacion(unittest.TestCase):
    def test_valores_directos(self):
        self.assertEqual(normalizar_clasificacion("PRODUCTO"), "PRODUCTO")
        self.assertEqual(normalizar_clasificacion(" producto "), "PRODUCTO")
        self.assertEqual(normalizar_clasificacion("Tania"), "TANIA")

    def test_alias_correcciones(self):
        for valor in (
            "LMS_correcciones",
            "LMS Correcciones",
            "lms-corrección",
            "LMSCORRECIONES",
            "Correcciones",
            "correccion",
            "LMS_CORRECC_2025",
            "LMSCORRECCIONES",
        ):
            with self.subTest(valor=valor):
                self.assertEqual(normalizar_clasificacion(valor), "LMS_CORRECCIONES")

    def test_invalidos(self):
        for valor in (None, "", "   ", "OTRO", "PRODUCTOS", 123, "LMS_Carga"):
            with self.subTest(valor=valor):
                self.assertIsNone(normalizar_clasificacion(valor))

    def test_clasificaciones_mapean_a_gcp(self):
        self.assertEqual(CLASIFICACIONES["PRODUCTO"], ("PRODUCTO", "LMS_Carga"))
        self.assertEqual(CLASIFICACIONES["TANIA"], ("TANIA", "LMS_Carga"))
        self.assertEqual(CLASIFICACIONES["LMS_CORRECCIONES"], ("PRODUCTO", "LMS_Carga"))

    def test_norm_text(self):
        self.assertEqual(excel.norm_text("  Ingeniería de Sistemas - 2025 "), "INGENIERIA_DE_SISTEMAS_2025")
        self.assertEqual(excel.norm_text(None), "")


class TestLote(unittest.TestCase):
    def test_clave(self):
        lote = Lote(2, "x", "PRODUCTO", "PRODUCTO", "PRODUCTO", "LMS_Carga", "ORI", "RAIZ", "ORI", "RAIZ")
        self.assertEqual(lote.clave, "ORI|RAIZ")

    def test_es_inmutable(self):
        lote = Lote(2, "x", "PRODUCTO", "PRODUCTO", "PRODUCTO", "LMS_Carga", "ORI", "RAIZ", "ORI", "RAIZ")
        with self.assertRaises(Exception):
            lote.etiqueta = "otra"  # type: ignore[misc]


class TestExcelInvalido(unittest.TestCase):
    def test_es_error_flujo_con_paso_excel(self):
        e = ExcelInvalido("motivo.", "accion.")
        self.assertIsInstance(e, ErrorFlujo)
        self.assertEqual(e.paso, "excel")
        self.assertEqual(str(e), "motivo. accion.")


class TestLeerLotesOk(BaseExcel):
    def test_cuatro_columnas_urls_ids_y_alias_cliente(self):
        ruta = self.excel(
            {
                1: [" CLIENTE ", "Etiqueta ", " ORIGEN", "Destino "],
                2: ["PRODUCTO", "Bogotá 01", URL_A, URL_OPEN],
                3: ["LMS_correcciones", "Medellín 02", URL_B, ID_RAIZ],
                4: ["Correcciones", "Cali 03", " 1EeOrigenCali_05 ", f'"{ID_RAIZ}"'],
                5: ["tania", "Pereira 04", "1FfOrigenPereira-06", URL_OPEN],
            }
        )
        lotes = leer_lotes(ruta)
        self.assertEqual([l.fila for l in lotes], [2, 3, 4, 5])
        self.assertEqual([l.etiqueta for l in lotes], ["Bogotá 01", "Medellín 02", "Cali 03", "Pereira 04"])
        self.assertEqual(
            [l.origen_id for l in lotes],
            ["1AaOrigenBogota_01", "1BbOrigenMedellin-02", "1EeOrigenCali_05", "1FfOrigenPereira-06"],
        )
        self.assertEqual(
            [l.destino_raiz_id for l in lotes],
            ["1CcDestinoRaiz_03", ID_RAIZ, ID_RAIZ, "1CcDestinoRaiz_03"],
        )
        self.assertEqual(
            [(l.clasificacion, l.cliente_gcp, l.raiz_gcp) for l in lotes],
            [
                ("PRODUCTO", "PRODUCTO", "LMS_Carga"),
                ("LMS_CORRECCIONES", "PRODUCTO", "LMS_Carga"),
                ("LMS_CORRECCIONES", "PRODUCTO", "LMS_Carga"),
                ("TANIA", "TANIA", "LMS_Carga"),
            ],
        )
        self.assertEqual([l.cliente_excel for l in lotes], ["PRODUCTO", "LMS_correcciones", "Correcciones", "tania"])
        self.assertEqual(lotes[0].origen_raw, URL_A)
        self.assertEqual(lotes[0].destino_raw, URL_OPEN)
        self.assertEqual(lotes[0].clave, "1AaOrigenBogota_01|1CcDestinoRaiz_03")

    def test_cabecera_desplazada_con_titulo_encima(self):
        ruta = self.excel(
            {
                1: ["RUTAS FÁBRICA DE CONTENIDO", None, None, None],
                3: CABECERA,
                4: ["PRODUCTO", "Bogotá", URL_A, ID_RAIZ],
            }
        )
        lotes = leer_lotes(ruta)
        self.assertEqual(len(lotes), 1)
        self.assertEqual(lotes[0].fila, 4)
        self.assertEqual(lotes[0].etiqueta, "Bogotá")

    def test_filas_vacias_intermedias(self):
        ruta = self.excel(
            {
                1: CABECERA,
                2: ["PRODUCTO", "Uno", URL_A, ID_RAIZ],
                5: ["TANIA", "Dos", URL_B, ID_RAIZ],
                6: [None, "", None, None],
                8: ["PRODUCTO", "Tres", "1GgOtroOrigen_07", URL_OPEN],
            }
        )
        lotes = leer_lotes(ruta)
        self.assertEqual([(l.fila, l.etiqueta) for l in lotes], [(2, "Uno"), (5, "Dos"), (8, "Tres")])

    def test_columnas_alias_tipo_y_programa(self):
        ruta = self.excel(
            {
                1: ["Tipo", "Programa", "Origen", "Destino", "Observaciones"],
                2: ["PRODUCTO", "Ingeniería", URL_A, ID_RAIZ, "nota"],
            }
        )
        lotes = leer_lotes(ruta)
        self.assertEqual(lotes[0].etiqueta, "Ingeniería")
        self.assertEqual(lotes[0].clasificacion, "PRODUCTO")

    def test_columna_clasificacion_con_acento(self):
        ruta = self.excel({1: ["Clasificación", "Etiqueta", "Origen", "Destino"], 2: ["TANIA", "x", URL_A, ID_RAIZ]})
        self.assertEqual(leer_lotes(ruta)[0].clasificacion, "TANIA")

    def test_etiqueta_vacia_y_sin_columna_etiqueta(self):
        con_columna = self.excel({1: CABECERA, 2: ["PRODUCTO", None, URL_A, ID_RAIZ]})
        self.assertEqual(leer_lotes(con_columna)[0].etiqueta, "Fila 2")
        sin_columna = self.excel({1: ["cliente", "origen", "destino"], 3: ["PRODUCTO", URL_A, ID_RAIZ]}, nombre="B.xlsx")
        self.assertEqual(leer_lotes(sin_columna)[0].etiqueta, "Fila 3")

    def test_etiqueta_numerica_se_convierte_a_texto(self):
        ruta = self.excel({1: CABECERA, 2: ["PRODUCTO", 2025, URL_A, ID_RAIZ]})
        self.assertEqual(leer_lotes(ruta)[0].etiqueta, "2025")

    def test_fila_con_solo_columna_extra_se_ignora(self):
        ruta = self.excel(
            {
                1: CABECERA + ["nota"],
                2: ["PRODUCTO", "Uno", URL_A, ID_RAIZ, ""],
                3: [None, None, None, None, "solo una observación"],
            }
        )
        self.assertEqual(len(leer_lotes(ruta)), 1)

    def test_datos_en_hoja_no_activa(self):
        libro = openpyxl.Workbook()
        libro.active.title = "Notas"
        libro.active.cell(row=1, column=1, value="Cualquier cosa")
        hoja = libro.create_sheet("Rutas")
        for c, v in enumerate(CABECERA, start=1):
            hoja.cell(row=1, column=c, value=v)
        for c, v in enumerate(["PRODUCTO", "x", URL_A, ID_RAIZ], start=1):
            hoja.cell(row=2, column=c, value=v)
        libro.active = 0
        ruta = self.dir / "dos_hojas.xlsx"
        libro.save(ruta)
        libro.close()
        self.assertEqual(len(leer_lotes(ruta)), 1)

    def test_acepta_ruta_como_texto(self):
        ruta = self.excel({1: CABECERA, 2: ["PRODUCTO", "x", URL_A, ID_RAIZ]})
        self.assertEqual(len(leer_lotes(str(ruta))), 1)


class TestLeerLotesErrores(BaseExcel):
    def test_errores_acumulados_en_una_sola_excepcion(self):
        ruta = self.excel(
            {
                1: CABECERA,
                2: [None, "Sin cliente", URL_A, ID_RAIZ],
                3: ["PRODUCTO", "Origen malo", "esto no es un enlace", ID_RAIZ],
                4: ["TANIA", "Sin destino", URL_B, None],
                5: ["PRODUCTO", "Bien", "1HhOrigenBueno_08", ID_RAIZ],
            }
        )
        e = self.leer_error(ruta)
        self.assertIn("Fila 2: cliente vacío", e.motivo)
        self.assertIn("Fila 3: enlace de origen inválido", e.motivo)
        self.assertIn("Fila 4: enlace de destino vacío", e.motivo)
        self.assertNotIn("Fila 5", e.motivo)
        self.assertIn("3 fila(s)", e.motivo)
        self.assertEqual(e.accion, "Corrige el Excel y vuelve a ejecutar.")
        self.assertEqual(e.paso, "excel")
        self.assertEqual(e.contexto, "RUTAS.xlsx")

    def test_varios_problemas_en_la_misma_fila(self):
        ruta = self.excel({1: CABECERA, 2: ["OTRO", "x", "https://ejemplo.com/otra", None]})
        e = self.leer_error(ruta)
        self.assertIn("cliente «OTRO» no reconocido", e.motivo)
        self.assertIn("enlace de origen inválido", e.motivo)
        self.assertIn("enlace de destino vacío", e.motivo)
        self.assertEqual(e.motivo.count("Fila 2:"), 1)

    def test_origen_igual_a_destino(self):
        ruta = self.excel({1: CABECERA, 2: ["PRODUCTO", "x", URL_A, "1AaOrigenBogota_01"]})
        e = self.leer_error(ruta)
        self.assertIn("Fila 2: origen y destino son la misma carpeta", e.motivo)

    def test_clave_repetida(self):
        ruta = self.excel(
            {
                1: CABECERA,
                2: ["PRODUCTO", "Uno", URL_A, ID_RAIZ],
                3: ["TANIA", "Dos", URL_B, ID_RAIZ],
                4: ["TANIA", "Repetido", "1AaOrigenBogota_01", f"https://drive.google.com/drive/folders/{ID_RAIZ}"],
            }
        )
        e = self.leer_error(ruta)
        self.assertIn("Fila 4: repite el origen y destino de la fila 2", e.motivo)
        self.assertNotIn("Fila 3", e.motivo)

    def test_falta_columna_cliente(self):
        ruta = self.excel({1: ["etiqueta", "origen", "destino"], 2: ["x", URL_A, ID_RAIZ]})
        e = self.leer_error(ruta)
        self.assertIn("falta la columna cliente", e.motivo)

    def test_faltan_varias_columnas(self):
        ruta = self.excel({1: ["cliente", "etiqueta"], 2: ["PRODUCTO", "x"]})
        e = self.leer_error(ruta)
        self.assertIn("origen, destino", e.motivo)

    def test_sin_cabecera_en_primeras_diez_filas(self):
        filas = {n: ["dato", n] for n in range(1, 12)}
        filas[12] = CABECERA
        e = self.leer_error(self.excel(filas))
        self.assertIn("cabecera", e.motivo)
        self.assertIn("10 filas", e.motivo)

    def test_excel_sin_lotes(self):
        e = self.leer_error(self.excel({1: CABECERA}))
        self.assertIn("no tiene lotes", e.motivo)
        vacio = self.excel({1: CABECERA, 2: [None, None, None, None]}, nombre="vacio.xlsx")
        self.assertIn("no tiene lotes", self.leer_error(vacio).motivo)

    def test_archivo_inexistente(self):
        e = self.leer_error(self.dir / "no_existe.xlsx")
        self.assertIn("No se encontró el Excel", e.motivo)
        self.assertEqual(e.paso, "excel")

    def test_archivo_que_no_es_excel(self):
        ruta = self.dir / "RUTAS.xlsx"
        ruta.write_text("esto no es un xlsx", encoding="utf-8")
        e = self.leer_error(ruta)
        self.assertIn("No se pudo leer el Excel RUTAS.xlsx", e.motivo)
        self.assertTrue(e.detalle)

    def test_excel_abierto_permission_error(self):
        ruta = self.excel({1: CABECERA, 2: ["PRODUCTO", "x", URL_A, ID_RAIZ]})
        with mock.patch("openpyxl.load_workbook", side_effect=PermissionError(13, "Permission denied")):
            e = self.leer_error(ruta)
        self.assertEqual(e.motivo, "El Excel RUTAS.xlsx está abierto o bloqueado.")
        self.assertEqual(e.accion, "Ciérralo y vuelve a ejecutar.")
        self.assertIn("está abierto", str(e))
        self.assertIn("PermissionError", e.detalle)


class TestResolverExcel(BaseExcel):
    def test_ruta_explicita_inexistente(self):
        with self.assertRaises(ExcelInvalido) as cm:
            resolver_excel(self.dir / "no_existe.xlsx")
        self.assertIn("No se encontró el Excel", cm.exception.motivo)
        self.assertIn("--excel", cm.exception.accion)
        self.assertEqual(cm.exception.paso, "excel")

    def test_ruta_explicita_existente(self):
        ruta = self.excel({1: CABECERA, 2: ["PRODUCTO", "x", URL_A, ID_RAIZ]})
        self.assertEqual(resolver_excel(ruta), ruta.resolve())

    def test_sin_ruta_usa_variable_de_entorno(self):
        ruta = self.excel({1: CABECERA, 2: ["PRODUCTO", "x", URL_A, ID_RAIZ]}, nombre="MIS_RUTAS.xlsx")
        with mock.patch.dict("os.environ", {"RUTAS_XLSX": str(ruta)}):
            self.assertEqual(resolver_excel(None), ruta.resolve())

    def test_sin_ruta_y_sin_excel_en_ninguna_parte(self):
        with mock.patch.dict("os.environ", {"RUTAS_XLSX": str(self.dir / "nada.xlsx")}), mock.patch(
            "flujo_lib.excel.resolver_rutas_excel", side_effect=FileNotFoundError("nada")
        ):
            with self.assertRaises(ExcelInvalido) as cm:
                resolver_excel(None)
        self.assertEqual(cm.exception.motivo, "No se encontró el archivo RUTAS.xlsx.")
        self.assertIn("--excel", cm.exception.accion)


if __name__ == "__main__":
    unittest.main()
