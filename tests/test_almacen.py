"""
Pruebas de flujo_lib.almacen (carpeta local y Cloud Storage) y de EstadoCorrida
usándolo. El cliente de Cloud Storage es un doble en memoria: no se toca la red.
"""

from __future__ import annotations

import io, json, os, tempfile, unittest
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import load_workbook

from flujo_lib.almacen import AlmacenGCS, AlmacenLocal, crear_almacen
from flujo_lib.estado import EstadoCorrida


# ---------------------------------------------------------------------------
# Doble de google.cloud.storage
# ---------------------------------------------------------------------------
class NoEncontrado(Exception):
    """Como google.api_core.exceptions.NotFound: objeto que no está en el bucket."""

    code = 404


class BlobFalso:
    def __init__(self, bucket: "BucketFalso", nombre: str):
        self.bucket = bucket
        self.name = nombre

    def download_as_bytes(self) -> bytes:
        try:
            return self.bucket.objetos[self.name][0]
        except KeyError:
            raise NoEncontrado(f"No existe {self.name}") from None

    def upload_from_string(self, datos, content_type=None) -> None:
        self.bucket.objetos[self.name] = (bytes(datos), content_type)

    def delete(self) -> None:
        try:
            del self.bucket.objetos[self.name]
        except KeyError:
            raise NoEncontrado(f"No existe {self.name}") from None


class BucketFalso:
    def __init__(self, nombre: str):
        self.name = nombre
        self.objetos: dict[str, tuple[bytes, str | None]] = {}

    def blob(self, nombre: str) -> BlobFalso:
        return BlobFalso(self, nombre)


class ClienteFalso:
    """Cliente en memoria con el trozo de API que usa AlmacenGCS."""

    def __init__(self):
        self.buckets: dict[str, BucketFalso] = {}

    def bucket(self, nombre: str) -> BucketFalso:
        return self.buckets.setdefault(nombre, BucketFalso(nombre))

    def list_blobs(self, bucket, prefix: str = ""):
        nombre = bucket if isinstance(bucket, str) else bucket.name
        objetos = self.buckets.setdefault(nombre, BucketFalso(nombre)).objetos
        return [SimpleNamespace(name=n) for n in sorted(objetos) if n.startswith(prefix or "")]


class TestAlmacenLocal(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.carpeta = Path(self._tmp.name) / "corridas"
        self.almacen = AlmacenLocal(self.carpeta)

    def tearDown(self):
        self._tmp.cleanup()

    def test_escribir_y_leer(self):
        self.almacen.escribir("RUTAS.estado.json", '{"ñ": "áé"}'.encode("utf-8"))
        self.assertEqual(self.almacen.leer("RUTAS.estado.json").decode("utf-8"), '{"ñ": "áé"}')
        self.assertTrue((self.carpeta / "RUTAS.estado.json").is_file())

    def test_leer_lo_que_no_existe_es_none(self):
        self.assertIsNone(self.almacen.leer("no_existe.json"))

    def test_sobrescribir(self):
        self.almacen.escribir("a.txt", b"viejo")
        self.almacen.escribir("a.txt", b"nuevo")
        self.assertEqual(self.almacen.leer("a.txt"), b"nuevo")
        self.assertEqual(self.almacen.listar(), ["a.txt"])

    def test_listar_con_prefijo_y_subcarpetas(self):
        self.almacen.escribir("RUTAS.estado.json", b"{}")
        self.almacen.escribir("OTRO.estado.json", b"{}")
        self.almacen.escribir("logs/20260922.log", b"linea")
        self.assertEqual(
            self.almacen.listar(), ["OTRO.estado.json", "RUTAS.estado.json", "logs/20260922.log"]
        )
        self.assertEqual(self.almacen.listar("RUTAS"), ["RUTAS.estado.json"])
        self.assertEqual(self.almacen.listar("logs/"), ["logs/20260922.log"])
        self.assertEqual(self.almacen.listar("nada"), [])

    def test_listar_carpeta_inexistente(self):
        self.assertEqual(AlmacenLocal(self.carpeta / "no" / "existe").listar(), [])

    def test_borrar_es_idempotente(self):
        self.almacen.escribir("a.txt", b"x")
        self.almacen.borrar("a.txt")
        self.assertIsNone(self.almacen.leer("a.txt"))
        self.almacen.borrar("a.txt")  # otra vez: no falla

    def test_escritura_atomica_sin_dejar_temporales(self):
        self.almacen.escribir("a.txt", b"x" * 1000)
        self.assertEqual([p.name for p in self.carpeta.iterdir()], ["a.txt"])

    def test_si_falla_el_reemplazo_no_queda_temporal_ni_se_pierde_lo_anterior(self):
        self.almacen.escribir("a.txt", b"viejo")
        with mock.patch.object(os, "replace", side_effect=OSError("disco lleno")):
            with self.assertRaises(OSError):
                self.almacen.escribir("a.txt", b"nuevo")
        self.assertEqual([p.name for p in self.carpeta.iterdir()], ["a.txt"])
        self.assertEqual(self.almacen.leer("a.txt"), b"viejo")

    def test_ruta_visible(self):
        self.assertEqual(
            self.almacen.ruta_visible("RUTAS.estado.json"),
            str(self.carpeta / "RUTAS.estado.json"),
        )


class TestAlmacenGCS(unittest.TestCase):
    def setUp(self):
        self.cliente = ClienteFalso()
        self.almacen = AlmacenGCS("fabrica-cun", "corridas", cliente=self.cliente)

    def _objetos(self) -> dict:
        return self.cliente.bucket("fabrica-cun").objetos

    def test_escribir_pone_el_prefijo_y_el_tipo(self):
        self.almacen.escribir("RUTAS.estado.json", b"{}")
        self.assertEqual(list(self._objetos()), ["corridas/RUTAS.estado.json"])
        self.assertEqual(self._objetos()["corridas/RUTAS.estado.json"][1], "application/json")

    def test_leer_y_sobrescribir(self):
        self.almacen.escribir("a.txt", b"uno")
        self.assertEqual(self.almacen.leer("a.txt"), b"uno")
        self.almacen.escribir("a.txt", b"dos")
        self.assertEqual(self.almacen.leer("a.txt"), b"dos")

    def test_leer_lo_que_no_existe_es_none(self):
        self.assertIsNone(self.almacen.leer("no_existe.json"))

    def test_un_error_que_no_sea_404_sube(self):
        with mock.patch.object(BlobFalso, "download_as_bytes", side_effect=RuntimeError("500")):
            with self.assertRaises(RuntimeError):
                self.almacen.leer("a.txt")

    def test_listar_quita_el_prefijo(self):
        self.almacen.escribir("RUTAS.estado.json", b"{}")
        self.almacen.escribir("logs/20260922.log", b"x")
        self._objetos()["otra_cosa/fuera.txt"] = (b"x", None)
        self.assertEqual(self.almacen.listar(), ["RUTAS.estado.json", "logs/20260922.log"])
        self.assertEqual(self.almacen.listar("logs/"), ["logs/20260922.log"])

    def test_borrar_es_idempotente(self):
        self.almacen.escribir("a.txt", b"x")
        self.almacen.borrar("a.txt")
        self.assertEqual(self._objetos(), {})
        self.almacen.borrar("a.txt")  # NotFound: no falla

    def test_ruta_visible(self):
        self.assertEqual(
            self.almacen.ruta_visible("RUTAS.estado.json"),
            "gs://fabrica-cun/corridas/RUTAS.estado.json",
        )

    def test_sin_prefijo(self):
        almacen = AlmacenGCS("fabrica-cun", cliente=self.cliente)
        almacen.escribir("a.txt", b"x")
        self.assertEqual(list(self._objetos()), ["a.txt"])
        self.assertEqual(almacen.listar(), ["a.txt"])
        self.assertEqual(almacen.ruta_visible("a.txt"), "gs://fabrica-cun/a.txt")

    def test_el_cliente_se_crea_tarde(self):
        """Sin usarlo no se crea ningún cliente de Google (ni se piden credenciales)."""
        almacen = AlmacenGCS("fabrica-cun")
        self.assertIsNone(almacen._cliente)
        with mock.patch("google.cloud.storage.Client", return_value=self.cliente) as crear:
            almacen.escribir("a.txt", b"x")
            almacen.escribir("b.txt", b"y")
        crear.assert_called_once()

    def test_un_error_de_borrado_que_no_sea_404_sube(self):
        with mock.patch.object(BlobFalso, "delete", side_effect=RuntimeError("500")):
            with self.assertRaises(RuntimeError):
                self.almacen.borrar("a.txt")


class TestCrearAlmacen(unittest.TestCase):
    def test_gs_da_almacen_de_cloud_storage(self):
        cliente = ClienteFalso()
        almacen = crear_almacen("gs://fabrica-cun/corridas/2026", cliente_gcs=cliente)
        self.assertIsInstance(almacen, AlmacenGCS)
        self.assertEqual((almacen.bucket, almacen.prefijo), ("fabrica-cun", "corridas/2026"))
        self.assertIs(almacen.cliente, cliente)

    def test_gs_sin_prefijo(self):
        almacen = crear_almacen("gs://fabrica-cun", cliente_gcs=ClienteFalso())
        self.assertEqual((almacen.bucket, almacen.prefijo), ("fabrica-cun", ""))

    def test_gs_sin_bucket_es_error(self):
        with self.assertRaises(ValueError):
            crear_almacen("gs://")

    def test_una_ruta_da_almacen_local(self):
        for destino in ("corridas", Path("corridas"), r"C:\Dev\corridas"):
            almacen = crear_almacen(destino)
            self.assertIsInstance(almacen, AlmacenLocal)
            self.assertEqual(almacen.carpeta, Path(str(destino)))


class Reloj:
    def __init__(self, inicio: datetime):
        self.actual = inicio

    def __call__(self) -> datetime:
        return self.actual

    def avanzar(self, segundos: float) -> None:
        self.actual = self.actual + timedelta(seconds=segundos)


def _lote(clave: str = "ORIG1|RAIZ1"):
    origen, raiz = clave.split("|")
    return SimpleNamespace(
        clave=clave, etiqueta="Lote", fila=2, cliente_excel="PRODUCTO",
        origen_id=origen, destino_raiz_id=raiz,
    )


class TestEstadoConAlmacen(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name) / "corridas"
        self.excel = Path(self._tmp.name) / "RUTAS.xlsx"
        self.reloj = Reloj(datetime(2026, 9, 22, 8, 0, 0))
        self.cliente = ClienteFalso()
        self.almacen = AlmacenGCS("fabrica-cun", "corridas", cliente=self.cliente)
        self.estado = EstadoCorrida.abrir(
            self.excel, self.dir, ahora=self.reloj, almacen=self.almacen
        )

    def tearDown(self):
        self._tmp.cleanup()

    def test_guarda_en_el_almacen_y_no_toca_el_disco(self):
        self.assertFalse(self.dir.exists())
        datos = json.loads(self.almacen.leer("RUTAS.estado.json").decode("utf-8"))
        self.assertEqual(datos["excel"], str(self.excel))
        self.assertEqual(datos["creado"], "2026-09-22T08:00:00")
        self.assertEqual(
            self.estado.ruta_visible_json, "gs://fabrica-cun/corridas/RUTAS.estado.json"
        )

    def test_reabrir_con_el_mismo_almacen_recupera_todo(self):
        lote = _lote()
        self.estado.iniciar_corrida({"schema": "fabrica"})
        self.estado.registrar_lote(lote)
        self.estado.set_destino(lote.clave, "DEST1", "Carpeta")
        self.estado.marcar(lote.clave, "destino", "ok")
        self.estado.avanzar(lote.clave, "clonacion", {"hechos": 4, "total": 4, "mensaje": "fin"})

        otro = EstadoCorrida.abrir(self.excel, self.dir, ahora=self.reloj, almacen=self.almacen)
        self.assertEqual(otro.datos, self.estado.datos)
        self.assertTrue(otro.completado(lote.clave, "destino"))
        self.assertEqual(otro.destino_de(lote.clave), ("DEST1", "Carpeta"))
        self.assertEqual(otro.progreso(lote.clave, "clonacion")["porcentaje"], 100)
        self.assertFalse(self.dir.exists())

    def test_exportar_excel_escribe_los_bytes_en_el_almacen(self):
        self.estado.iniciar_corrida({})
        self.estado.registrar_lote(_lote())
        ruta = self.estado.exportar_excel()
        self.assertEqual(ruta, self.estado.ruta_excel_estado)
        self.assertFalse(self.dir.exists())  # nada en el disco
        datos = self.almacen.leer("RUTAS.estado.xlsx")
        self.assertIsNotNone(datos)
        wb = load_workbook(io.BytesIO(datos))
        self.assertEqual(wb.sheetnames, ["Estado", "Corridas"])
        self.assertEqual(wb["Estado"].cell(2, 1).value, "Lote")
        self.assertEqual(
            self.estado.ruta_visible_excel, "gs://fabrica-cun/corridas/RUTAS.estado.xlsx"
        )

    def test_almacen_local_inyectado_equivale_al_comportamiento_de_siempre(self):
        carpeta = Path(self._tmp.name) / "otras_corridas"
        estado = EstadoCorrida.abrir(
            self.excel, self.dir, ahora=self.reloj, almacen=crear_almacen(carpeta)
        )
        estado.registrar_lote(_lote())
        estado.marcar("ORIG1|RAIZ1", "destino", "ok")
        guardado = json.loads((carpeta / "RUTAS.estado.json").read_text(encoding="utf-8"))
        self.assertEqual(guardado["lotes"]["ORIG1|RAIZ1"]["pasos"]["destino"]["estado"], "ok")
        self.assertEqual([p.name for p in carpeta.iterdir()], ["RUTAS.estado.json"])

    def test_json_danado_en_el_almacen_da_error_flujo(self):
        from flujo_lib.mensajes import ErrorFlujo

        self.almacen.escribir("RUTAS.estado.json", b"{esto no es json")
        with self.assertRaises(ErrorFlujo) as ctx:
            EstadoCorrida.abrir(self.excel, self.dir, ahora=self.reloj, almacen=self.almacen)
        self.assertIn("está dañado", ctx.exception.motivo)


if __name__ == "__main__":
    unittest.main()
